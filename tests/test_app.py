from datetime import timedelta
from decimal import Decimal
from io import BytesIO
import pytest
from easy import create_app,today
from easy.models import db,User,Order

@pytest.fixture
def app():
    app=create_app({'TESTING':True,'SECRET_KEY':'test-secret-'*5,'SQLALCHEMY_DATABASE_URI':'sqlite://','WTF_CSRF_ENABLED':False,'RATELIMIT_ENABLED':False,'APP_ENV':'development'})
    with app.app_context():
        db.create_all()
        for username,role in [('admin','admin'),('team','equipe')]:
            u=User(username=username,name=username,role=role);u.set_password('strong-password-123');db.session.add(u)
        db.session.commit()
    yield app
    with app.app_context():db.drop_all()
@pytest.fixture
def client(app):
    c=app.test_client();c.post('/login',data={'username':'admin','password':'strong-password-123'});return c

def create_order(client,app):
    assert client.post('/clientes/novo',data={'name':'Cliente teste','active':'y'}).status_code==302
    assert client.post('/veiculos/novo',data={'client_id':1,'plate':'abc-1d23','model':'Ducato','active':'y'}).status_code==302
    r=client.post('/ordens/nova',data={'number':'','vehicle_id':1,'entry_date':today().isoformat(),'due_date':today().isoformat(),'mileage':12345,'responsible':'Diogo','problem':'Não refrigera','discount':'0'})
    assert r.status_code==302
    with app.app_context():
        order=db.session.scalar(db.select(Order));assert order.number.startswith('OS-');return order.id

def version(app,ident):
    with app.app_context():return str(db.session.get(Order,ident).version)

def add_item(client,app,ident,kind='servico',description='Diagnóstico',qty='1',price='150',done=False):
    return client.post(f'/ordens/{ident}/itens',data={'version':version(app,ident),'kind':kind,'description':description,'quantity':qty,'unit_price':price,'done':'y' if done else ''})

def test_complete_workflow(client,app):
    ident=create_order(client,app)
    assert add_item(client,app,ident,done=True).status_code==302
    assert add_item(client,app,ident,'peca','Filtro',qty='2',price='80').status_code==302
    with app.app_context():
        order=db.session.get(Order,ident);assert order.total==Decimal('310.00');assert len(order.checks)==9
    client.post(f'/ordens/{ident}/status',data={'version':version(app,ident),'status':'Entregue','delivered_date':today().isoformat()})
    with app.app_context():assert db.session.get(Order,ident).status!='Entregue'
    for stage in ['entrada','saida']:
        with app.app_context():ids=[c.id for c in db.session.get(Order,ident).checks if c.stage==stage]
        assert client.post(f'/ordens/{ident}/checklist',data={'version':version(app,ident),'stage':stage,**{f'result_{i}':'ok' for i in ids}}).status_code==302
    assert client.post(f'/ordens/{ident}/status',data={'version':version(app,ident),'status':'Entregue','delivered_date':today().isoformat()}).status_code==302
    with app.app_context():
        order=db.session.get(Order,ident);assert order.status=='Entregue';assert order.delivered_date==today();assert len(order.events)>=6
    assert add_item(client,app,ident).status_code==409
    for path in ['/', '/clientes','/veiculos','/veiculos/1','/ordens','/ordens/nova',f'/ordens/{ident}','/relatorios','/usuarios',f'/ordens/{ident}/imprimir']:
        assert client.get(path).status_code==200,path
    from openpyxl import load_workbook
    workbook=load_workbook(BytesIO(client.get('/relatorios/exportar').data))
    assert workbook.sheetnames==['Ordens de serviço','Serviços e peças','Checklist']
    assert workbook.active['L2'].value==310

def test_stale_updates_and_money(client,app):
    ident=create_order(client,app);v=version(app,ident)
    add_item(client,app,ident,price='0.10',qty='3')
    assert client.post(f'/ordens/{ident}/itens',data={'version':v,'kind':'peca','description':'Old','quantity':'1','unit_price':'9'}).status_code==409
    with app.app_context():assert db.session.get(Order,ident).total==Decimal('0.30')
    add_item(client,app,ident,price='-1')
    with app.app_context():assert len(db.session.get(Order,ident).items)==1

def test_permissions_csrf_and_login(app):
    client=app.test_client()
    assert client.get('/clientes').status_code==302
    assert client.get('/relatorios/exportar').status_code==302
    assert client.post('/login?next=https://evil.example',data={'username':'team','password':'strong-password-123'}).location=='/'
    assert client.get('/usuarios').status_code==403
    app.config['WTF_CSRF_ENABLED']=True
    assert client.post('/clientes/novo',data={'name':'Unauthorized'}).status_code==400
    assert client.get('/').headers['Cache-Control']=='no-store'

def test_lockout_and_session_revocation(app):
    client=app.test_client()
    for _ in range(5):client.post('/login',data={'username':'admin','password':'wrong'})
    assert client.post('/login',data={'username':'admin','password':'strong-password-123'}).status_code==200
    with app.app_context():
        u=db.session.get(User,1);u.locked_until=None;db.session.commit()
    assert client.post('/login',data={'username':'admin','password':'strong-password-123'}).status_code==302
    with app.app_context():
        u=db.session.get(User,1);u.session_version+=1;db.session.commit()
    assert client.get('/').status_code==302

def test_discount_and_historical_client(client,app):
    ident=create_order(client,app);add_item(client,app,ident,price='200')
    payload={'version':version(app,ident),'number':'CUSTOM-001','vehicle_id':1,'entry_date':today().isoformat(),'due_date':today().isoformat(),'mileage':'12','responsible':'Diogo','problem':'Teste','discount':'25'}
    assert client.post(f'/ordens/{ident}/editar',data=payload).status_code==302
    with app.app_context():
        order=db.session.get(Order,ident);assert order.total==Decimal('175');item_id=order.items[0].id
    client.post(f'/ordens/{ident}/itens/{item_id}/acao',data={'version':version(app,ident),'action':'remove'})
    with app.app_context():assert len(db.session.get(Order,ident).items)==1
    client.post('/clientes/novo',data={'name':'Novo dono','active':'y'})
    client.post('/veiculos/1/editar',data={'client_id':2,'plate':'ABC1D23','model':'Ducato','active':'y'})
    with app.app_context():assert db.session.get(Order,ident).client_id==1

def test_reports_filters_and_formula_injection(client,app):
    ident=create_order(client,app);add_item(client,app,ident,description='=2+2',price='99')
    from openpyxl import load_workbook
    book=load_workbook(BytesIO(client.get('/relatorios/exportar').data))
    assert book['Serviços e peças']['C2'].data_type!='f'
    assert 'R$ 0,00' in client.get('/relatorios?status=Entregue').data.decode()
    assert client.get('/relatorios?start=invalid').status_code==400

def test_future_delivery_rejected(client,app):
    ident=create_order(client,app);add_item(client,app,ident,done=True)
    for stage in ['entrada','saida']:
        with app.app_context():ids=[c.id for c in db.session.get(Order,ident).checks if c.stage==stage]
        client.post(f'/ordens/{ident}/checklist',data={'version':version(app,ident),'stage':stage,**{f'result_{i}':'ok' for i in ids}})
    client.post(f'/ordens/{ident}/status',data={'version':version(app,ident),'status':'Entregue','delivered_date':(today()+timedelta(days=1)).isoformat()})
    with app.app_context():assert db.session.get(Order,ident).status!='Entregue'

def test_private_photo_upload_and_login_gate(client,app,monkeypatch):
    from PIL import Image
    import cloudinary.uploader
    import requests
    ident=create_order(client,app)
    app.config['CLOUDINARY_URL']='cloudinary://123:test-secret@example'
    uploaded={}
    def upload(stream,**kwargs):
        uploaded.update(kwargs)
        return {'public_id':'easy-servicos/os-1/test','format':'png'}
    monkeypatch.setattr(cloudinary.uploader,'upload',upload)
    data=BytesIO();Image.new('RGB',(16,16)).save(data,format='PNG');data.seek(0)
    r=client.post(f'/ordens/{ident}/fotos',data={'version':version(app,ident),'stage':'entrada','caption':'Entrada','file':(data,'foto.png')},content_type='multipart/form-data')
    assert r.status_code==302;assert uploaded['type']=='authenticated'
    assert app.test_client().get('/fotos/1').status_code==302
    class Response:
        content=b'photo';headers={'Content-Type':'image/png'}
        def raise_for_status(self):pass
    monkeypatch.setattr(requests,'get',lambda *a,**kw:Response())
    assert client.get('/fotos/1').data==b'photo'
    assert client.get('/fotos/1').headers['Cache-Control']=='no-store'

def test_changes_revoke_ready_status(client,app):
    ident=create_order(client,app);add_item(client,app,ident,done=True)
    with app.app_context():
        order=db.session.get(Order,ident)
        for check in order.checks:check.result='ok'
        db.session.commit()
    client.post(f'/ordens/{ident}/status',data={'version':version(app,ident),'status':'Pronto para entrega'})
    with app.app_context():assert db.session.get(Order,ident).status=='Pronto para entrega'
    add_item(client,app,ident,'servico','Novo serviço',done=False)
    with app.app_context():assert db.session.get(Order,ident).status=='Em serviço'

def test_pwa_assets_and_camera(client,app):
    import json
    manifest=json.loads(client.get('/static/manifest.webmanifest').data)
    assert manifest['start_url']=='/' and manifest['display']=='standalone'
    for icon in manifest['icons']:assert client.get(icon['src']).status_code==200
    worker=client.get('/sw.js');assert worker.status_code==200
    assert worker.headers['Service-Worker-Allowed']=='/'
    assert client.get('/static/offline.html').status_code==200
    assert client.get('/instalar').status_code==200
    assert app.test_client().get('/instalar').status_code==302
    ident=create_order(client,app)
    app.config['CLOUDINARY_URL']='cloudinary://123:test-secret@example'
    assert 'capture="environment"' in client.get(f'/ordens/{ident}').data.decode()

def test_login_with_real_csrf_protection(app):
    import re
    app.config['WTF_CSRF_ENABLED']=True
    client=app.test_client()
    def token(path):
        page=client.get(path)
        found=re.search(r'name="csrf_token"[^>]*value="([^"]+)"',page.data.decode())
        assert found,path
        return found.group(1)
    csrf=token('/login')
    response=client.post('/login',data={'username':'admin','password':'strong-password-123','csrf_token':csrf})
    assert response.status_code==302
    assert client.get('/').status_code==200
    csrf=token('/clientes/novo')
    assert client.post('/clientes/novo',data={'name':'Cliente CSRF real','active':'y','csrf_token':csrf}).status_code==302
    assert client.post('/clientes/novo',data={'name':'Sem token'}).status_code==400
    assert client.post('/logout',data={'csrf_token':csrf}).status_code==302
    assert client.get('/').status_code==302

def test_production_database_and_cookie_configuration():
    with pytest.raises(RuntimeError,match='Produção exige'):
        create_app({'SECRET_KEY':'x'*48,'APP_ENV':'production','SQLALCHEMY_DATABASE_URI':'sqlite://','RATELIMIT_ENABLED':False})
    production=create_app({'SECRET_KEY':'x'*48,'APP_ENV':'production','SQLALCHEMY_DATABASE_URI':'postgresql+psycopg2://test:test@localhost/test','RATELIMIT_ENABLED':False})
    assert production.config['SESSION_COOKIE_SECURE'] is True
    assert production.config['SESSION_COOKIE_HTTPONLY'] is True
    assert production.config['SQLALCHEMY_ENGINE_OPTIONS']['connect_args']['sslmode']=='require'
    assert production.config['WTF_CSRF_TIME_LIMIT']==28800

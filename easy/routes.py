import csv
import io
import re
from datetime import datetime, timedelta, date
from decimal import Decimal, ROUND_HALF_UP
from functools import wraps
from urllib.parse import urlsplit, unquote
from flask import Blueprint, render_template, request, redirect, url_for, flash, abort, session, send_file, current_app
from flask_login import current_user, login_user, logout_user, login_required
from sqlalchemy import or_, func, case
from sqlalchemy.orm import selectinload
from werkzeug.security import generate_password_hash, check_password_hash
from . import db, limiter, today
from .models import User, Client, Vehicle, Order, Item, Check, Event, Photo, STATUSES, ENTRY_CHECKS, EXIT_CHECKS, utcnow
from .forms import LoginForm, ClientForm, OrderForm, ItemForm, StatusForm, PhotoForm, UserForm, FlowForm, ServiceForm

web=Blueprint('web',__name__)
DUMMY_HASH=generate_password_hash('dummy-password-never-a-login')

def admin_required(fn):
    @wraps(fn)
    @login_required
    def wrapped(*args,**kwargs):
        if current_user.role!='admin':abort(403)
        return fn(*args,**kwargs)
    return wrapped

def get_order(ident,**kwargs):
    order=db.get_or_404(Order,ident,**kwargs)
    if order.deleted_at:abort(404)
    return order

def event(order,message,invalidate_ready=False):
    stamp=utcnow()
    db.session.add(Event(order=order,user_id=current_user.id,created_at=stamp,description=message))
    order.updated_at=stamp

def mutable(order):
    submitted=request.form.get('version')
    if submitted is None or submitted!=str(order.version):
        abort(409,description='A OS foi alterada. Atualize a página e tente novamente.')

def keep_discount_valid(order,subtotal):
    if order.discount>subtotal:
        previous=order.discount;order.discount=max(subtotal,Decimal('0'))
        event(order,f'Desconto histórico ajustado de {previous} para {order.discount} após alteração dos itens.')

def order_choices(form,order=None):
    condition=Client.active.is_(True) & Client.deleted_at.is_(None)
    if order:condition=or_(condition,Client.id==order.client_id)
    clients=db.session.scalars(db.select(Client).where(condition).order_by(Client.name)).all()
    form.client_id.choices=[(client.id,client.name) for client in clients]

def populate_order(form,order):
    for name in ['client_id','entry_date','due_date','mileage','responsible','problem','diagnosis','notes']:
        value=getattr(form,name).data
        if name in ['responsible','problem','diagnosis','notes']:value=(value or '').strip()
        setattr(order,name,value)
    order.vehicle_plate=re.sub(r'[\s-]','',(form.plate.data or '').strip().upper())
    order.vehicle_model=(form.model.data or '').strip()
    order.vehicle_fleet=(form.fleet.data or '').strip()
    # A relação antiga fica interna para compatibilidade; nunca altera outra OS.
    if order.vehicle_plate:
        vehicle=db.session.scalar(db.select(Vehicle).where(Vehicle.plate==order.vehicle_plate))
        if vehicle is None:
            vehicle=Vehicle(plate=order.vehicle_plate,model=order.vehicle_model,
                            fleet=order.vehicle_fleet,client_id=form.client_id.data)
            db.session.add(vehicle)
        order.vehicle=vehicle
    else:
        order.vehicle=None


def validate_order(form,order=None):
    if form.due_date.data and form.entry_date.data and form.due_date.data<form.entry_date.data:
        form.due_date.errors.append('A entrega prevista deve ser igual ou posterior à entrada.');return False
    client=db.session.get(Client,form.client_id.data)
    if not client or (not client.active and (not order or order.client_id!=client.id)):
        form.client_id.errors.append('Selecione um cliente ativo.');return False
    plate=re.sub(r'[\s-]','',(form.plate.data or '').strip().upper())
    if len(plate)>10:
        form.plate.errors.append('Informe uma identificação de até 10 caracteres.');return False
    number=(form.number.data or '').strip().upper()
    existing=db.session.scalar(db.select(Order).where(Order.number==number)) if number else None
    if existing and (not order or existing.id!=order.id):
        form.number.errors.append('Este número de OS já está em uso.');return False
    return True

@web.route('/login',methods=['GET','POST'])
@limiter.limit('10 per minute')
def login():
    if current_user.is_authenticated:return redirect(url_for('web.dashboard'))
    form=LoginForm()
    if form.validate_on_submit():
        username=form.username.data.strip().lower()
        user=db.session.scalar(db.select(User).where(User.username==username).with_for_update())
        valid=user.check_password(form.password.data) if user else check_password_hash(DUMMY_HASH,form.password.data)
        if user and user.active and valid and (not user.locked_until or user.locked_until<=utcnow()):
            user.failures=0;user.locked_until=None;db.session.commit()
            session.clear();login_user(user,remember=False);session.permanent=True
            target=request.args.get('next','')
            parsed=urlsplit(target)
            return redirect(target if target.startswith('/') and not target.startswith('//') and not parsed.netloc and not parsed.scheme and '\\' not in target else url_for('web.dashboard'))
        if user and (not user.locked_until or user.locked_until<=utcnow()):
            user.failures+=1
            if user.failures>=5:user.locked_until=utcnow()+timedelta(minutes=15);user.failures=0
            db.session.commit()
        else:db.session.rollback()
        flash('Usuário ou senha inválidos, ou acesso temporariamente bloqueado.','error')
    return render_template('login.html',form=form)

@web.post('/logout')
@login_required
def logout():
    logout_user();session.clear();return redirect(url_for('web.login'))

@web.get('/')
@login_required
def dashboard():
    active=Order.deleted_at.is_(None) & ~Order.status.in_(['Entregue','Faturado','Recebido','Cancelada'])
    pending_finance=Order.deleted_at.is_(None) & Order.status.in_(['Entregue','Faturado']) & Order.received_date.is_(None)
    day=today()
    totals=db.session.execute(db.select(
        func.count(case((active,1))),
        func.count(case((active & (Order.due_date==day),1))),
        func.count(case((active & (Order.due_date<day),1))),
        func.count(case((pending_finance,1))))).one()
    counts=dict(zip(['open','today','late'],totals[:3]))
    finance_count=totals[3]
    orders=db.session.scalars(db.select(Order).where(or_(active,pending_finance)).options(db.joinedload(Order.vehicle),db.joinedload(Order.client),selectinload(Order.items)).order_by(active.desc(),Order.due_date.asc().nulls_last(),Order.id).limit(12)).all()
    return render_template('dashboard.html',orders=orders,counts=counts,finance_count=finance_count)

@web.get('/clientes')
@login_required
def clients():
    q=request.args.get('q','').strip()[:150]
    query=db.select(Client).where(Client.deleted_at.is_(None))
    if q:query=query.where(Client.name.ilike('%'+q+'%'))
    page=db.paginate(query.order_by(Client.name),per_page=30,error_out=False)
    return render_template('clients.html',page=page,q=q)

@web.route('/clientes/novo',methods=['GET','POST'])
@web.route('/clientes/<int:ident>/editar',methods=['GET','POST'])
@login_required
def client_form(ident=None):
    obj=db.get_or_404(Client,ident) if ident else Client()
    if ident and obj.deleted_at:abort(404)
    form=ClientForm(obj=obj if ident else None)
    if form.validate_on_submit():
        form.populate_obj(obj);obj.name=obj.name.strip();db.session.add(obj);db.session.commit()
        flash('Cliente salvo.','success');return redirect(url_for('web.clients'))
    return render_template('form.html',form=form,title='Editar cliente' if ident else 'Novo cliente',back=url_for('web.clients'))

@web.post('/clientes/<int:ident>/excluir')
@admin_required
def client_delete(ident):
    client=db.get_or_404(Client,ident)
    if client.deleted_at:abort(404)
    client.deleted_at=utcnow()
    client.active=False
    db.session.commit()
    flash('Cliente excluído. As OS existentes foram preservadas.','success')
    return redirect(url_for('web.clients'))

# Links antigos continuam válidos, mas não há cadastro separado de veículos.
@web.get('/veiculos')
@login_required
def vehicles():
    return redirect(url_for('web.orders'))

@web.route('/veiculos/novo',methods=['GET','POST'])
@web.route('/veiculos/<int:ident>/editar',methods=['GET','POST'])
@login_required
def vehicle_form(ident=None):
    return redirect(url_for('web.order_new'),code=303 if request.method=='POST' else 302)

@web.get('/veiculos/<int:ident>')
@login_required
def vehicle_history(ident):
    vehicle=db.get_or_404(Vehicle,ident)
    return redirect(url_for('web.orders',q=vehicle.plate))

def filtered_orders():
    query=db.select(Order).where(Order.deleted_at.is_(None)).outerjoin(Vehicle).join(Client,Order.client_id==Client.id).options(db.joinedload(Order.vehicle),db.joinedload(Order.client),selectinload(Order.items))
    q=request.args.get('q','').strip()[:150];status=request.args.get('status','')
    if q:query=query.where(or_(Order.number.ilike('%'+q+'%'),func.coalesce(Order.vehicle_plate,Vehicle.plate).ilike('%'+q+'%'),func.coalesce(Order.vehicle_fleet,Vehicle.fleet).ilike('%'+q+'%'),Client.name.ilike('%'+q+'%')))
    if status in STATUSES:query=query.where(Order.status==status)
    if request.args.get('open')=='1':query=query.where(~Order.status.in_(['Entregue','Faturado','Recebido','Cancelada']))
    if request.args.get('finance')=='1':query=query.where(Order.status.in_(['Entregue','Faturado']),Order.received_date.is_(None))
    if request.args.get('due')=='today':query=query.where(Order.due_date==today(),~Order.status.in_(['Entregue','Faturado','Recebido','Cancelada']))
    if request.args.get('due')=='late':query=query.where(Order.due_date<today(),~Order.status.in_(['Entregue','Faturado','Recebido','Cancelada']))
    for key,operator in [('start',lambda d:Order.entry_date>=d),('end',lambda d:Order.entry_date<=d)]:
        raw=request.args.get(key)
        if raw:
            try:query=query.where(operator(date.fromisoformat(raw)))
            except ValueError:abort(400,description='Data de filtro inválida.')
    return query.order_by(Order.entry_date.desc(),Order.id.desc())

@web.get('/ordens')
@login_required
def orders():
    page=db.paginate(filtered_orders(),per_page=25,error_out=False)
    return render_template('orders.html',page=page,title='Ordens de serviço',q=request.args.get('q',''),status=request.args.get('status',''),vehicle=None)

@web.route('/ordens/nova',methods=['GET','POST'])
@login_required
def order_new():
    form=OrderForm()
    order_choices(form)
    if form.validate_on_submit() and validate_order(form):
        order=Order();populate_order(form,order);order.entered_at=utcnow();order.status_changed_at=order.entered_at
        if not order.entry_date:order.entry_date=today()
        order.number=(form.number.data or '').strip().upper() or None
        db.session.add(order);db.session.flush()
        if not order.number:order.number=f'OS-{order.created_at.year}-{order.id:06d}'
        for stage,labels in [('entrada',ENTRY_CHECKS),('saida',EXIT_CHECKS)]:
            for label in labels:db.session.add(Check(order=order,stage=stage,label=label))
        event(order,'OS criada.');db.session.commit()
        flash('OS criada. Agora adicione os serviços e as peças.','success')
        return redirect(url_for('web.order_detail',ident=order.id))
    return render_template('order_form.html',form=form,title='Nova ordem de serviço',back=url_for('web.orders'),empty=not form.client_id.choices,empty_message='Cadastre um cliente ativo para abrir uma OS. O veículo é preenchido aqui.')

@web.route('/ordens/<int:ident>/editar',methods=['GET','POST'])
@login_required
def order_edit(ident):
    order=get_order(ident)
    form=OrderForm(obj=order);order_choices(form,order)
    if request.method=='GET':form.version.data=str(order.version)
    if form.validate_on_submit():
        mutable(order)
        if validate_order(form,order):
            number=(form.number.data or order.number).strip().upper();populate_order(form,order);order.number=number
            event(order,'Dados da OS atualizados.');db.session.commit();flash('OS atualizada.','success')
            return redirect(url_for('web.order_detail',ident=ident))
    return render_template('order_form.html',form=form,title=f'Editar {order.number}',back=url_for('web.order_detail',ident=ident))

@web.post('/ordens/<int:ident>/excluir')
@admin_required
def order_delete(ident):
    order=get_order(ident)
    if request.form.get('version')!=str(order.version):
        abort(409,description='A OS foi alterada. Atualize a página antes de excluir.')
    order.deleted_at=utcnow()
    event(order,'OS excluída pelo administrador.',invalidate_ready=False)
    db.session.commit()
    flash('OS excluída.','success')
    return redirect(url_for('web.orders'))

@web.get('/ordens/<int:ident>')
@login_required
def order_detail(ident):
    order=get_order(ident,options=[db.joinedload(Order.vehicle),db.joinedload(Order.client),selectinload(Order.items),selectinload(Order.checks),selectinload(Order.photos),selectinload(Order.events).joinedload(Event.user)])
    status_form=StatusForm(status=order.status,version=str(order.version),delivered_date=order.delivered_date or today());status_form.status.choices=[(x,x) for x in STATUSES]
    flow_form=FlowForm(version=str(order.version),action_date=today())
    if current_user.role=='admin':
        flow_form.action.choices += [('clear_approval','Corrigir: remover aprovação'),('clear_invoice','Corrigir: remover faturamento'),('clear_received','Corrigir: remover recebimento')]
    return render_template('order.html',order=order,status_form=status_form,flow_form=flow_form,item_form=ItemForm(version=str(order.version)),service_form=ServiceForm(version=str(order.version)),photo_form=PhotoForm(version=str(order.version)),photos_enabled=bool(current_app.config['CLOUDINARY_URL']))

def record_status(order,target,day=None,reference=''):
    stamp=utcnow();day=day or today();previous=order.status
    fields={'Aprovado':'approved','Em serviço':'started','Pronto para entrega':'ready',
            'Entregue':'delivered','Faturado':'invoiced','Recebido':'received'}
    if target=='Aguardando avaliação':order.entered_at=stamp
    prefix=fields.get(target)
    if prefix:
        setattr(order,prefix+'_date',day);setattr(order,prefix+'_at',stamp)
    # Preserve other movement records. A selection never clears another step.
    order.status=target;order.status_changed_at=stamp
    event(order,f'Status: {previous} → {target}.'+(f' Observação: {reference}' if reference else ''))

@web.post('/ordens/<int:ident>/fluxo')
@login_required
def order_flow(ident):
    order=get_order(ident);form=FlowForm()
    if current_user.role=='admin':
        form.action.choices += [('clear_approval','Remover aprovação'),('clear_invoice','Remover faturamento'),('clear_received','Remover recebimento')]
    if not form.validate_on_submit():abort(400,description='Confira a etapa selecionada.')
    mutable(order)
    action=form.action.data;reference=(form.reference.data or '').strip()
    if action.startswith('clear_'):
        if current_user.role!='admin':abort(403)
        field={'clear_approval':'approved','clear_invoice':'invoiced','clear_received':'received'}[action]
        setattr(order,field+'_date',None);setattr(order,field+'_at',None)
        event(order,f'Registro {field} removido.'+(f' Observação: {reference}' if reference else ''))
    else:
        target={'approve':'Aprovado','invoice':'Faturado','receive':'Recebido'}[action]
        record_status(order,target,form.action_date.data,reference)
    db.session.commit();flash('Fluxo atualizado.','success')
    return redirect(url_for('web.order_detail',ident=ident))

@web.post('/ordens/<int:ident>/status')
@login_required
def order_status(ident):
    order=get_order(ident);form=StatusForm();form.status.choices=[(x,x) for x in STATUSES]
    if not form.validate_on_submit():abort(400,description='Confira o status selecionado.')
    mutable(order)
    target=form.status.data
    day=form.delivered_date.data if target=='Entregue' else None
    record_status(order,target,day)
    db.session.commit();flash('Fluxo atualizado.','success')
    return redirect(url_for('web.order_detail',ident=ident))

@web.post('/ordens/<int:ident>/checklist')
@login_required
def checklist(ident):
    order=get_order(ident);mutable(order)
    stage=request.form.get('stage')
    if stage not in ['entrada','saida']:abort(400)
    for check in order.checks:
        if check.stage!=stage:continue
        result=request.form.get(f'result_{check.id}','pendente');notes=request.form.get(f'notes_{check.id}','').strip()
        if result not in (['pendente','ok','avaria','na'] if stage=='entrada' else ['pendente','ok']) or len(notes)>500:abort(400)
        check.result=result;check.notes=notes
    event(order,f'Checklist de {"entrada" if stage=="entrada" else "entrega"} atualizado.');db.session.commit()
    flash('Checklist salvo.','success');return redirect(url_for('web.order_detail',ident=ident))

@web.post('/ordens/<int:ident>/servicos')
@login_required
def service_add(ident):
    order=get_order(ident);mutable(order);form=ServiceForm()
    if not form.validate_on_submit():
        abort(400,description='Confira o tipo de serviço, peças utilizadas e profissional responsável.')
    service_type=form.service_type.data
    title={'refrigeracao':'Refrigeração','estrutura':'Estrutura'}[service_type]
    item=Item(order=order,kind='servico',service_type=service_type,description=title,
              used_parts=(form.used_parts.data or '').strip(),responsible=(form.responsible.data or '').strip(),
              quantity=1,unit_price=Decimal('0'),done=False)
    db.session.add(item);event(order,f'Serviço {title} cadastrado.');db.session.commit()
    flash('Serviço salvo.','success');return redirect(url_for('web.order_detail',ident=ident))

@web.post('/ordens/<int:ident>/itens')
@login_required
def item_add(ident):
    order=get_order(ident);mutable(order);form=ItemForm()
    if not form.validate_on_submit():
        flash('Confira descrição, quantidade e valor do item.','error');return redirect(url_for('web.order_detail',ident=ident))
    item=Item(order=order);form.populate_obj(item)
    event(order,f'Adicionado {"serviço" if item.kind=="servico" else "peça"}: {item.description}.');db.session.add(item);db.session.commit()
    flash('Item adicionado.','success');return redirect(url_for('web.order_detail',ident=ident))

@web.route('/ordens/<int:ident>/itens/<int:item_id>/editar',methods=['GET','POST'])
@login_required
def item_edit(ident,item_id):
    order=get_order(ident);item=db.get_or_404(Item,item_id)
    if item.order_id!=ident:abort(404)
    if item.service_type:
        form=ServiceForm(obj=item)
        if request.method=='GET':form.version.data=str(order.version)
        if form.validate_on_submit():
            mutable(order)
            item.service_type=form.service_type.data
            item.description={'refrigeracao':'Refrigeração','estrutura':'Estrutura'}[form.service_type.data]
            item.used_parts=(form.used_parts.data or '').strip()
            item.responsible=(form.responsible.data or '').strip()
            event(order,f'Serviço atualizado: {item.description}.');db.session.commit()
            flash('Serviço salvo.','success');return redirect(url_for('web.order_detail',ident=ident))
        return render_template('form.html',form=form,title='Editar serviço',back=url_for('web.order_detail',ident=ident))
    form=ItemForm(obj=item)
    if request.method=='GET':form.version.data=str(order.version)
    if form.validate_on_submit():
        mutable(order)
        proposed=(form.quantity.data*form.unit_price.data).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP)
        keep_discount_valid(order,order.subtotal()-item.total+proposed)
        form.populate_obj(item);event(order,f'Item atualizado: {item.description}.');db.session.commit()
        flash('Item atualizado.','success');return redirect(url_for('web.order_detail',ident=ident))
    return render_template('form.html',form=form,title='Editar serviço / peça',back=url_for('web.order_detail',ident=ident))

@web.post('/ordens/<int:ident>/itens/<int:item_id>/acao')
@login_required
def item_action(ident,item_id):
    order=get_order(ident);mutable(order);item=db.session.get(Item,item_id)
    if item is None:abort(404,description='Este serviço já foi excluído ou não está mais disponível. Atualize a OS.')
    if item.order_id!=ident:abort(404)
    action=request.form.get('action')
    if action=='remove':
        keep_discount_valid(order,order.subtotal()-item.total)
        event(order,f'Item removido: {item.description}.');db.session.delete(item)
    elif action=='toggle':
        item.done=not item.done;event(order,f'Item {"concluído" if item.done else "reaberto"}: {item.description}.')
    else:abort(400)
    db.session.commit();return redirect(url_for('web.order_detail',ident=ident))

def configure_photos():
    import cloudinary
    parsed=urlsplit(current_app.config['CLOUDINARY_URL'])
    if parsed.scheme!='cloudinary' or not parsed.hostname or not parsed.username or not parsed.password:
        raise RuntimeError('CLOUDINARY_URL inválida.')
    cloudinary.config(cloud_name=parsed.hostname,api_key=unquote(parsed.username),api_secret=unquote(parsed.password),secure=True)

@web.post('/ordens/<int:ident>/fotos')
@login_required
@limiter.limit('20 per hour')
def photo_add(ident):
    order=get_order(ident);mutable(order)
    if not current_app.config['CLOUDINARY_URL']:
        flash('Configure o armazenamento de fotos para habilitar os anexos.','error');return redirect(url_for('web.order_detail',ident=ident))
    form=PhotoForm()
    if not form.validate_on_submit():abort(400)
    from PIL import Image, UnidentifiedImageError
    import cloudinary
    import cloudinary.uploader
    file=form.file.data
    try:
        img=Image.open(file.stream)
        if img.format not in ['JPEG','PNG','WEBP'] or img.width*img.height>25000000:raise ValueError()
        img.verify();file.stream.seek(0)
    except (UnidentifiedImageError,OSError,ValueError,Image.DecompressionBombError):
        flash('Envie uma imagem JPEG, PNG ou WebP válida, com até 25 megapixels.','error');return redirect(url_for('web.order_detail',ident=ident))
    configure_photos()
    try:
        uploaded=cloudinary.uploader.upload(file.stream,folder=f'easy-servicos/os-{ident}',type='authenticated',resource_type='image',allowed_formats=['jpg','jpeg','png','webp'],timeout=30)
    except Exception:
        current_app.logger.warning('Falha no envio da foto para o armazenamento externo.')
        flash('Não foi possível enviar a foto. Tente novamente.','error');return redirect(url_for('web.order_detail',ident=ident))
    db.session.add(Photo(order=order,public_id=uploaded['public_id'],format=uploaded['format'],stage=form.stage.data,caption=form.caption.data))
    event(order,f'Foto de {form.stage.data} anexada.');
    try:db.session.commit()
    except Exception:
        db.session.rollback()
        try:cloudinary.uploader.destroy(uploaded['public_id'],type='authenticated',resource_type='image')
        except Exception:current_app.logger.warning('Falha ao limpar foto órfã.')
        raise
    flash('Foto anexada.','success');return redirect(url_for('web.order_detail',ident=ident))

@web.get('/fotos/<int:ident>')
@login_required
def photo_view(ident):
    photo=db.get_or_404(Photo,ident)
    if photo.order.deleted_at:abort(404)
    if not current_app.config['CLOUDINARY_URL']:abort(404)
    import cloudinary
    import cloudinary.utils
    import requests
    configure_photos()
    url=cloudinary.utils.private_download_url(photo.public_id,photo.format,resource_type='image',type='authenticated',expires_at=int(datetime.now().timestamp())+60,attachment=False)
    try:
        response=requests.get(url,timeout=20)
        response.raise_for_status()
        mime=response.headers.get('Content-Type','').split(';')[0]
        if mime not in ['image/jpeg','image/png','image/webp']:abort(502)
        return send_file(io.BytesIO(response.content),mimetype=mime)
    except requests.RequestException:abort(502)

@web.get('/relatorios')
@login_required
def reports():
    page=db.paginate(filtered_orders(),per_page=30,error_out=False)
    # Aggregate all filtered records, not just the current page.
    ids=filtered_orders().with_only_columns(Order.id).order_by(None).subquery()
    item_total=db.session.execute(db.select(Item.kind,func.sum(func.round(Item.quantity*Item.unit_price,2))).where(Item.order_id.in_(db.select(ids.c.id))).group_by(Item.kind)).all()
    totals={kind:Decimal(str(value)).quantize(Decimal('0.01')) for kind,value in item_total};discount=db.session.scalar(db.select(func.sum(Order.discount)).where(Order.id.in_(db.select(ids.c.id)))) or Decimal('0')
    labor=totals.get('servico',Decimal('0'));parts=totals.get('peca',Decimal('0'))
    return render_template('reports.html',page=page,labor=labor,parts=parts,discount=discount,total=labor+parts-discount,q=request.args.get('q',''),status=request.args.get('status',''))

def safe_cell(value):
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):return "'"+value
    return value

@web.get('/relatorios/exportar')
@login_required
def report_export():
    from openpyxl import Workbook
    rows=db.session.scalars(filtered_orders().options(selectinload(Order.checks)).limit(10001)).all()
    if len(rows)>10000:abort(400,description='Filtre o período para exportar até 10 mil OS.')
    book=Workbook();sheet=book.active;sheet.title='Ordens de serviço'
    sheet.append(['OS','Cliente','Placa','Frota','Entrada','Previsão entrega','Entrega real','Status','Serviços','Peças','Desconto','Total','Aprovação','Início do serviço','Pronto','Faturamento','Recebimento','Etapa do fluxo','Registro da entrada','Registro da aprovação','Registro do início','Registro pronto','Registro da saída','Registro do faturamento','Registro do recebimento','Último movimento de status'])
    items=book.create_sheet('Serviços e peças');items.append(['OS','Tipo','Descrição','Quantidade','Valor unitário','Total','Responsável','Concluído','Tipo de serviço','Peças utilizadas'])
    checks=book.create_sheet('Checklist');checks.append(['OS','Etapa','Item','Resultado','Observações'])
    for order in rows:
        sheet.append([safe_cell(x) for x in [order.number,order.client.name,order.plate,order.fleet,order.entry_date,order.due_date,order.delivered_date,order.status,order.subtotal('servico'),order.subtotal('peca'),order.discount,order.total,order.approved_date,order.started_date,order.ready_date,order.invoiced_date,order.received_date,order.flow_label,*[current_app.jinja_env.filters['localtime'](stamp) if stamp else '' for stamp in [order.entered_at,order.approved_at,order.started_at,order.ready_at,order.delivered_at,order.invoiced_at,order.received_at,order.status_changed_at]]]])
        for item in order.items:items.append([safe_cell(x) for x in [order.number,item.kind,item.description,item.quantity,item.unit_price,item.total,item.responsible,'Sim' if item.done else 'Não',item.service_type or '',item.used_parts or '']])
        for check in order.checks:checks.append([safe_cell(x) for x in [order.number,check.stage,check.label,check.result,check.notes]])
    from openpyxl.styles import Font, PatternFill
    for ws in book:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='125B75')
        for col in ws.columns:
            width=min(55,max(16,max(len(str(c.value or '')) for c in col)+2));ws.column_dimensions[col[0].column_letter].width=width
            for cell in col[1:]:
                if isinstance(cell.value,date):cell.number_format='dd/mm/yyyy'
        money_cols=[9,10,11,12] if ws==sheet else [5,6] if ws==items else []
        for n in money_cols:
            for col in ws.iter_cols(min_col=n,max_col=n,min_row=2):
                for cell in col:cell.number_format='"R$" #,##0.00'
    output=io.BytesIO();book.save(output);output.seek(0)
    return send_file(output,as_attachment=True,download_name=f'relatorio-servicos-{today()}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@web.get('/ordens/<int:ident>/imprimir')
@login_required
def order_print(ident):
    return render_template('print.html',order=get_order(ident,options=[db.joinedload(Order.vehicle),db.joinedload(Order.client),selectinload(Order.items),selectinload(Order.checks)]))

@web.get('/usuarios')
@admin_required
def users():
    return render_template('users.html',users=db.session.scalars(db.select(User).order_by(User.name)).all())

@web.route('/usuarios/novo',methods=['GET','POST'])
@web.route('/usuarios/<int:ident>/editar',methods=['GET','POST'])
@admin_required
def user_form(ident=None):
    obj=db.get_or_404(User,ident) if ident else User()
    form=UserForm(obj=obj if ident else None)
    if not ident:form.password.label.text='Senha (mínimo 6 caracteres)'
    if form.validate_on_submit():
        username=form.username.data.strip().lower()
        if not ident and not form.password.data:form.password.errors.append('Defina a senha inicial.')
        elif db.session.scalar(db.select(User).where(User.username==username,User.id!=(ident or 0))):form.username.errors.append('Usuário já existe.')
        elif ident==current_user.id and (not form.active.data or form.role.data!='admin'):
            form.role.errors.append('Você não pode remover seu próprio acesso de administrador.')
        else:
            changed=ident and (obj.role!=form.role.data or obj.active!=form.active.data or bool(form.password.data))
            obj.username=username;obj.name=form.name.data.strip();obj.role=form.role.data;obj.active=form.active.data
            if form.password.data:obj.set_password(form.password.data);obj.failures=0;obj.locked_until=None
            if changed:obj.session_version+=1
            db.session.add(obj);db.session.commit()
            if ident==current_user.id and changed:login_user(obj)
            flash('Usuário salvo.','success');return redirect(url_for('web.users'))
    return render_template('form.html',form=form,title='Editar usuário' if ident else 'Novo usuário',back=url_for('web.users'))


@web.get('/sw.js')
@limiter.exempt
def service_worker():
    response=current_app.send_static_file('sw.js')
    response.headers['Service-Worker-Allowed']='/'
    response.headers['Cache-Control']='no-cache'
    return response

@web.get('/instalar')
@login_required
def install():
    return render_template('install.html')
